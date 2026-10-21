#!/usr/bin/env python3
"""Validates a generated tracker with openpyxl (structure) and, if available, LibreOffice (opens + recalculates).

Usage: python3 scripts/validate_workbook.py path/to/tracker.xlsx
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

EXPECTED_SHEETS = ["Gantt Chart", "WBS", "API List Pending", "Escalation Matrix", "Last Week Task", "This Week Task", "Finance"]
EXPECTED_TABLES = {
    "Gantt Chart": ["tblGantt"],
    "WBS": ["tblWBS"],
    "API List Pending": ["tblAPIs"],
    "Escalation Matrix": ["tblEscalationMatrix", "tblEscalationLog"],
    "Last Week Task": ["tblLastWeek"],
    "This Week Task": ["tblThisWeek"],
    "Finance": ["tblFinance"],
}


def main(path: str) -> int:
    errors = []
    wb = openpyxl.load_workbook(path)
    if wb.sheetnames != EXPECTED_SHEETS:
        errors.append(f"sheets {wb.sheetnames} != {EXPECTED_SHEETS}")
    for name, tables in EXPECTED_TABLES.items():
        ws = wb[name]
        if sorted(ws.tables.keys()) != sorted(tables):
            errors.append(f"{name}: tables {list(ws.tables.keys())}")
        if ws.freeze_panes is None:
            errors.append(f"{name}: no freeze panes")
        if name != "Escalation Matrix" and not ws.data_validations.dataValidation and name not in ("API List Pending", "Finance"):
            errors.append(f"{name}: no data validation")
    wbs = wb["WBS"]
    header = [c.value for c in wbs[5]]
    if header[:4] != ["WBS ID", "Task ID", "Phase", "Task"]:
        errors.append(f"WBS header {header}")
    if not str(wbs["K6"].value).startswith("="):
        errors.append(f"WBS Current Due is not a formula: {wbs['K6'].value!r}")
    fin = wb["Finance"]
    fin_header = [c.value for c in fin[5]]
    if fin_header[3:7] != ["Amount Raised and Overdue (for Go-Live)", "Amount that can be raised on immediate basis", "Total", "Amount to be raised"]:
        errors.append(f"Finance header {fin_header}")
    if not str(fin["F6"].value).startswith("="):
        errors.append(f"Finance Total is not a formula: {fin['F6'].value!r}")
    print(f"openpyxl: {len(wb.sheetnames)} sheets, WBS rows={wbs.max_row - 5}")

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        with tempfile.TemporaryDirectory() as tmp:
            res = subprocess.run(
                [soffice, "--headless", "--convert-to", "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,2",
                 "--outdir", tmp, path],
                capture_output=True, text=True, timeout=120,
            )
            csv = Path(tmp) / (Path(path).stem + "-WBS.csv")
            if res.returncode != 0 or not csv.exists():
                errors.append(f"LibreOffice failed: {res.stderr or res.stdout}")
            else:
                rows = [l.split(",") for l in csv.read_text().splitlines()]
                data = rows[5] if len(rows) > 5 else []
                print(f"LibreOffice: opened and recalculated; WBS row 6 Current Due={data[10] if len(data) > 10 else '?'}")
                if len(data) > 10 and (data[10] in ("", "#NAME?", "Err:508") or data[10].startswith("#")):
                    errors.append(f"Current Due formula did not evaluate: {data[10]!r}")
    else:
        print("LibreOffice not found; skipped recalculation check")

    for e in errors:
        print("ERROR:", e)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
